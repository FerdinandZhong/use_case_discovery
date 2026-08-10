from pptx import Presentation
F = "Use Case Discovery Deck (for AE) - ZH.pptx"
NEW = {
 "客户":"客户数据", "靶点识别":"交易数据", "药物开发":"生产数据", "内部数据":"供应链数据",
 "真实世界数据":"设备传感器", "体内试验":"文档与合同", "体外试验":"风控与合规", "临床试验":"运营日志",
}
def collapse(t): return t.replace("\x0b","").replace("\n","").replace(" ","").strip()
def walk(shapes):
    n=0
    for sh in shapes:
        if sh.shape_type==6: n+=walk(sh.shapes); continue
        if not sh.has_text_frame: continue
        key=collapse(sh.text_frame.text)
        if key in NEW:
            tf=sh.text_frame
            # keep only first paragraph
            for extra in list(tf.paragraphs)[1:]:
                extra._p.getparent().remove(extra._p)
            p0=tf.paragraphs[0]
            runs=list(p0.runs)
            for extra in runs[1:]:
                extra._r.getparent().remove(extra._r)
            runs[0].text = NEW[key]   # preserves white/bold/size formatting of run 0
            n+=1
    return n
prs=Presentation(F)
c=sum(walk(s.shapes) for s in prs.slides)
prs.save(F)
print("re-themed nodes:", c)
