#!/usr/bin/env python3
"""Translate the AE discovery deck to Simplified Chinese, in place, run by run.
Layout / shapes / fonts untouched. Matches on stripped run text; leaves
technical terms, product names, TBox/ABox, '+', URLs, and '[ ] ' unchanged."""
import sys
from pptx import Presentation

SRC = "Use Case Discovery Deck (for AE).pptx"
OUT = "Use Case Discovery Deck (for AE) - ZH.pptx"

# english (stripped)  \t  简体中文
TSV = """
AI USE CASE DISCOVERY\tAI 用例发现
Ready your data. Ground your AI.\t让数据就绪，让 AI 有据可依。
Let's find the use cases worth building.\t一起找出值得落地的用例。
Prepared for AEs and SEs\t面向 AE 与 SE
Skilled people burn hours on repetitive, manual steps that don’t scale.\t专业人才把大量时间耗在无法规模化的重复手工步骤上。
AI can take that toil, if it truly understands your business.\tAI 可以接手这些繁琐工作，前提是它真正理解你的业务。
Why AI, and why it needs your data\t为什么需要 AI，以及它为何离不开你的数据
Your experts spend their days on manual work\t你的专家每天都在做手工活
Out of the box it doesn’t:\t开箱即用时它做不到：
it fills the gaps by guessing.\t它只能靠猜测来填补空白。
Trust erosion:\t信任受损：
Confident, wrong answers erode trust, and the fix isn't a bigger model. So what's actually missing?\t自信却错误的回答会侵蚀信任，而解决之道并非更大的模型。那么，究竟缺了什么？
Action:\t行动：
Through a use case discovery workshop, understand your data, your platform, your daily operations, and get ready for AI solutions\t通过一场用例发现工作坊，摸清你的数据、平台与日常运营，为 AI 方案做好准备
The gap isn't capability — it's a missing semantic layer\t差距不在能力，而在缺失的语义层
Why bare RAG + bare Agent hallucinate\t为什么纯 RAG 与纯 Agent 会产生幻觉
Bare RAG\t纯 RAG
Matches text, not meaning.\t匹配的是文本，而非语义。
Finds a doc with “inventory,” but can't tell ERP-context from WMS-context — two different things.\t能找到含“库存”的文档，却分不清 ERP 语境与 WMS 语境，二者截然不同。
Vector similarity ≠ business semantics.\t向量相似 ≠ 业务语义。
Bare Agent\t纯 Agent
Tools return codes; the LLM guesses meaning.\t工具只返回代码，LLM 只能猜测其含义。
MATERIAL.ALLOY_CD / defect_cd=7 — the AI has no idea which alloy grade / weld defect it means, so it fills the gap.\tMATERIAL.ALLOY_CD / defect_cd=7：AI 根本不知道它指哪种合金牌号或焊接缺陷，只能凭空填补。
Tools return codes, not meaning.\t工具返回的是代码，不是语义。
The missing layer that lets AI read data by business meaning = the Ontology.\t让 AI 按业务语义读取数据的那一缺失层 = 本体（Ontology）。
A governed semantic layer, grounded in provenance.\t一个受治理、可溯源的语义层。
AI Context Layer\tAI 上下文层
Bring all your data into one platform\t把所有数据汇入同一平台
Databases, documents, images, logs, and comms in one place\t数据库、文档、图像、日志与沟通记录，集中一处
Solves the "where is my data?" problem\t解决“我的数据在哪里？”的问题
Knowledge Silos: the real challenge\t知识孤岛：真正的挑战
Data together ≠ Knowledge connected\t数据聚合 ≠ 知识连通
Missing semantic relationships and context\t缺少语义关系与上下文
Team still can't discover cross-domain insights\t团队仍无法发现跨领域洞见
The Knowledge Graph\t知识图谱
Entities + Relationships + Semantic Context = Connected Knowledge\t实体 + 关系 + 语义上下文 = 连通的知识
Turns a data lake into an intelligent knowledge network\t把数据湖变成智能知识网络
Becomes the context layer that makes AI smarter and grounded\t成为让 AI 更聪明、更有据可依的上下文层
Data together is not knowledge connected\t数据聚合不等于知识连通
From Data Driven  to Cognitive\t从数据驱动到认知驱动
Your question (plain language)\t你的问题（自然语言）
Map to your vocabulary\t映射到你的业务词汇
Pull the exact governed data\t拉取精确的受治理数据
Grounded answer\t有据可依的答案
TBox — semantic standard (no business data)\tTBox：语义标准（不含业务数据）
Class · Relation · Property · Axiom.\t类 · 关系 · 属性 · 公理。
Auto: Part / BOM / Variant / Supplier / ECN; “part_id globally unique”, “superseded rev not in active BOM”.\t自动：Part / BOM / Variant / Supplier / ECN；“part_id 全局唯一”“被取代的版本不出现在有效 BOM 中”。
ABox — instances (read on demand)\tABox：实例（按需读取）
Governed rows, read on demand.\t受治理的数据行，按需读取。
Declarative mapping registers only where-to-fetch & how-to-transform — reads no data, aligns semantics.\t声明式映射只登记“到哪取数”与“如何转换”，不读取数据，只对齐语义。
The ontology, in two parts\t本体的两个部分
Give AI your vocabulary: TBox + ABox\t把你的业务词汇交给 AI：TBox + ABox
Painful, repetitive work\t痛苦而重复的工作
Data we can ground\t可用于支撑的数据
A use case worth building\t值得落地的用例
Signs worth a look\t值得关注的信号
•   Repetitive cross-referencing across disconnected systems\t•   在互不相通的系统间反复交叉核对
•   Manual review, triage, and approvals that gate the work\t•   人工审核、分诊与审批卡住了流程
•   Answers that live in people’s heads, not any system\t•   答案只存在于人的脑子里，而非任何系统
What makes a strong first pick\t什么样的首选用例才算好
•   The work is painful and repeats often\t•   工作痛苦且频繁重复
•   The data to ground it exists (that’s your ontology)\t•   支撑它的数据已经存在（也就是你的本体）
•   There’s a clear decision or action at the end\t•   末端有明确的决策或行动
Where could AI help in your day-to-day?\tAI 能在你的日常工作中帮上什么忙？
Now your data can ground AI. So where do we point it first?\t现在你的数据已能支撑 AI，那我们先从哪里入手？
Proven agentic workflows on Cloudera\t在 Cloudera 上已验证的智能体（Agentic）工作流
Agents that replace manual review, grounded in governed data\t以受治理数据为支撑、替代人工审核的智能体
Pre Discovery Homework: Conduct Interview/Survey\t发现前准备：开展访谈 / 问卷调查
Owner : Field Specialist. Use this when the Customer /SE is not sure what problem to tackle\t负责人：现场专家。当客户 / SE 不确定要解决什么问题时使用
IDENTIFY THE FRICTION\t识别摩擦点
Focusing on the pain points where AI typically excels (Repetition, Data Overload, Poor  Prediction).\t聚焦 AI 通常最擅长的痛点（重复劳动、数据过载、预测不佳）。
The  Identification:\t识别：
Identify one core process in your department that consumes at least 20% of your team’s weekly manual effort.\t找出你部门中一个核心流程，它至少占用团队每周 20% 的手工工作量。
Examples: Manually categorizing emails, extracting data from PDFs, summarizing long reports, or verifying invoices.\t示例：手工分类邮件、从 PDF 中提取数据、总结长篇报告，或核对发票。
The Bottleneck:\t瓶颈：
In the process mentioned above, what is the primary "blocker"?\t在上述流程中，主要的“卡点”是什么？
[ ] Waiting for human approval/review.\t[ ] 等待人工审批 / 审核。
[ ] Accessing data from multiple disparate systems.\t[ ] 需从多个彼此独立的系统获取数据。
[ ] Cleaning or formatting inconsistent data.\t[ ] 清洗或整理不一致的数据。
[ ] Complex decision-making that requires specialized knowledge.\t[ ] 需要专业知识的复杂决策。
The "If I Only Knew" Question:\t“要是我能预知”之问：
If you could predict one business variable with 90% accuracy (e.g., customer churn, equipment failure, next month's demand), which one would change your strategy the most?\t如果你能以 90% 的准确率预测某一个业务变量（例如客户流失、设备故障、下月需求），哪一个最能改变你的策略？
IMPORTANT :\t重要：
Use open ended questions, and use GPT tools to get deeper in the customer domain\t多用开放式问题，并借助 GPT 工具深入了解客户所在领域
Data & Tooling Readiness\t数据与工具就绪度
Assess availability of the data,  controls and approvals needed\t评估所需数据、管控与审批的可获得性
Data Format:\t数据形态：
Where does the information for this process currently live? (Select all that apply)\t该流程的信息目前存放在哪里？（多选）
[ ] Structured Databases (SQL, Snowflake, ERP)\t[ ] 结构化数据库（SQL、Snowflake、ERP）
[ ] Unstructured Documents (PDFs, Word, Slide decks)\t[ ] 非结构化文档（PDF、Word、幻灯片）
[ ] Internal Communications (Slack, Emails, Teams)\t[ ] 内部沟通（Slack、邮件、Teams）
[ ] Tribal Knowledge (Only lives in people's heads)\t[ ] 隐性经验（只存在于人的脑子里）
Data Quality:\t数据质量：
On a scale of 1–5, how much do you trust the accuracy and cleanliness of this data today?\t以 1-5 分计，你当前有多信任这些数据的准确性与整洁度？
(1 = Total Mess, 5 = Ready for Analysis)\t（1 = 一团糟，5 = 可直接分析）
Data Quality and Readiness will be your #1 area where you could get quickly derailed\t数据质量与就绪度将是最容易让你迅速翻车的头号环节
Strategic Impact\t战略影响
Helps in prioritization across use cases\t帮助在多个用例之间排定优先级
The Value Driver:\t价值驱动：
If we automated or "augmented" this task with AI, what would be the primary win?\t如果用 AI 把这项任务自动化或“增强”，主要收益会是什么？
Cost Savings:\t成本节约：
Reducing headcount or operational spend.\t减少人力或运营开支。
Velocity:\t速度：
Doing the same task 10x faster.\t把同样的任务做快 10 倍。
Quality:\t质量：
Reducing human error and increasing consistency.\t减少人为错误，提升一致性。
Revenue:\t营收：
Creating a new product or better customer experience.\t打造新产品或更好的客户体验。
Risk Tolerance:\t风险容忍度：
How "accurate" does this AI need to be for you to trust it?\t这个 AI 要多“准”，你才会信任它？
Low Risk:\t低风险：
70% accuracy is fine (e.g., internal drafting).\t70% 准确率即可（例如内部起草）。
High Risk:\t高风险：
99% accuracy is required (e.g., financial reporting, legal compliance)\t需要 99% 准确率（例如财务报告、法律合规）
Identifying the Strategic Impact helps the customer also internally sell the project\t厘清战略影响，也有助于客户在内部推动该项目立项
Blue Sky Idea\t天马行空的设想
Smart Question to ask in case the customer has an “hidden agenda” and want their own “project” to be priortized\t当客户另有“小算盘”、希望优先推进自己“项目”时，可以巧问的问题
Blue Sky Thinking:\t天马行空：
If you had a "Magic Assistant" that knew everything about your Department/Domains historical data and could answer any question in seconds, what is the first question you would ask it?\t如果你有一位“神奇助手”，通晓你部门 / 领域的全部历史数据，并能在几秒内回答任何问题，你最先会问它什么？
Identifying the Strategic Impact helps the customer also internally sell the project,\t厘清战略影响，也有助于客户在内部推动该项目立项，
Also build your own questionnaire list based on domain or engage Industry Solutions team\t也可基于领域自建问题清单，或联系行业解决方案团队
We walk your daily workflows together: the steps, the data, and the tools.\t我们一起梳理你的日常工作流：步骤、数据与工具。
We pinpoint where AI agents can take the manual load off your team.\t我们精准找出 AI 智能体能为团队卸下手工负担的环节。
It takes about 10 minutes, and it seeds the workshop.\t大约只需 10 分钟，就能为工作坊打好基础。
Scan the code or open your link to start:\t扫码或打开你的链接即可开始：
Let’s map your use cases\t一起梳理你的用例
Your discovery deck\t你的用例发现问卷
Customers\t客户
Target ID\t靶点识别
Drug Dev\t药物开发
Internal Data\t内部数据
RWD\t真实世界数据
In Vivo Testing\t体内试验
In Vitro Testing\t体外试验
Clinical Trials\t临床试验
Drug\t药物
Dev\t开发
Internal\t内部
Data\t数据
Clinical\t临床
Trials\t试验
AI Specialist & FDE\tAI 专家与 FDE
Banking support\t银行客服支持
Chatbot with memory\t具备记忆的聊天机器人
Recalls:\t记住：
the customer across every session\t跨每一次会话记住客户
Retrieves:\t检索：
live account data from the lakehouse\t从数据湖读取实时账户数据
Reasons:\t推理：
over cross-session patterns to escalate\t跨会话模式推理以决定是否升级
Result:\t结果：
no more “I explained this last time”\t不再有“上次我已经解释过了”
Trade fraud\t贸易欺诈
Six-agent pipeline\t六智能体流水线
Screens:\t筛查：
every declaration, not a sample\t筛查每一份申报，而非抽样
Extracts:\t提取：
invoice fields via OCR\t通过 OCR 提取发票字段
Checks:\t核查：
prices, sanctions, collusion patterns\t价格、制裁名单与串通模式
30-90 min review → seconds\t审核从 30-90 分钟缩短到数秒
Manufacturing\t制造业
Defect triage + guardrails\t缺陷分诊 + 护栏
Grounds:\t依据：
every claim in an SOP § or data row\t每条结论都溯源到 SOP 条款或数据行
Blocks:\t拦截：
invented part numbers and action codes\t拦截编造的零件号与操作码
Abstains:\t弃答：
when the evidence is insufficient\t证据不足时选择不作答
human checks ~100% → <20%\t人工核查从约 100% 降至 <20%
"""

MAP = {}
for line in TSV.strip("\n").splitlines():
    if "\t" not in line:
        continue
    en, zh = line.split("\t", 1)
    MAP[en] = zh

prs = Presentation(SRC)
translated, misses = 0, []


def tr_paras(paragraphs):
    global translated
    for p in paragraphs:
        for r in p.runs:
            st = r.text.strip()
            if not st:
                continue
            if st in MAP:
                r.text = MAP[st]
                translated += 1
            elif len(st) > 5 and any(c.isascii() and c.isalpha() for c in st) and st not in ("TBox", "ABox"):
                if not st.startswith("http") and st not in ("[ ]",):
                    misses.append(st)


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP -> recurse
            walk(sh.shapes)
        elif sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    tr_paras(cell.text_frame.paragraphs)
        elif sh.has_text_frame:
            tr_paras(sh.text_frame.paragraphs)


for s in prs.slides:
    walk(s.shapes)

prs.save(OUT)
print(f"translated {translated} runs -> {OUT}")
if misses:
    print(f"\nPOSSIBLE MISSES ({len(misses)}):")
    for m in misses:
        print("  ", repr(m))
