#!/usr/bin/env python3
"""Build a per-customer discovery deck from the master template.

Fills {{CUSTOMER}} / {{DECK_URL}} and drops a QR for the discovery-deck URL onto
the "Let's map your use cases" slide.

    python build_deck.py --customer "Acme Corp" \
        --url "https://<host>/d/acme-1a2b3c" --out ~/Desktop/Acme-discovery.pptx
"""
import argparse
import io
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
DEFAULT_MASTER = HERE / "discovery-deck-master.pptx"
CTA_MARKER = "map your use cases"  # identifies the discovery CTA slide


def replace_placeholders(prs, mapping):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    for k, v in mapping.items():
                        if k in run.text:
                            run.text = run.text.replace(k, v)


def find_cta_slide(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and CTA_MARKER in shape.text_frame.text:
                return slide
    return None


def qr_png(url):
    buf = io.BytesIO()
    qrcode.make(url, border=1).save(buf, format="PNG")
    buf.seek(0)
    return buf


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--customer", required=True)
    ap.add_argument("--url", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--master", default=str(DEFAULT_MASTER))
    args = ap.parse_args()

    prs = Presentation(args.master)
    replace_placeholders(prs, {"{{CUSTOMER}}": args.customer, "{{DECK_URL}}": args.url})

    slide = find_cta_slide(prs)
    if slide is None:
        raise SystemExit("Could not find the CTA slide (marker missing)")
    # QR in the open right-center area (slide is 10 x 5.625 in)
    size = Inches(1.5)
    slide.shapes.add_picture(qr_png(args.url), Inches(7.7), Inches(2.9), height=size, width=size)

    out = Path(args.out).expanduser()
    prs.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
